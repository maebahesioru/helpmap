"""Multimodal ingestion — turn course material into source-linked content units.

Supported inputs:
  - PDF  (textbook chapters, papers, notes)   -> pymupdf, per-page text + embedded image refs
  - PPTX (lecture slide decks)                -> python-pptx, per-slide text + speaker notes
  - TXT/MD                                    -> paragraph chunks
  - VTT/SRT/TXT transcripts (lecture videos)  -> timestamped chunks

Every unit keeps a precise origin: file, page / slide / timestamp.
"""
from __future__ import annotations

import hashlib
import json
import re
from dataclasses import dataclass, asdict
from pathlib import Path


@dataclass
class Unit:
    id: str
    text: str
    source_type: str  # pdf | pptx | text | transcript
    source_file: str
    location: str  # "p.12" | "slide 4" | "00:12:30" | "para 3"
    page: int | None = None
    image_count: int = 0
    topics: list[str] | None = None

    def to_dict(self) -> dict:
        return asdict(self)


def _uid(file: str, loc: str, idx: int) -> str:
    h = hashlib.sha1(f"{file}:{loc}:{idx}".encode()).hexdigest()[:8]
    return f"u{h}"


def _chunk_paragraphs(text: str, max_chars: int = 900) -> list[str]:
    """Split text into paragraph-ish chunks under max_chars."""
    paras = [p.strip() for p in re.split(r"\n\s*\n|\n(?=[A-Z0-9])", text) if p.strip()]
    chunks, buf = [], ""
    for p in paras:
        if len(buf) + len(p) + 2 <= max_chars:
            buf = f"{buf}\n\n{p}" if buf else p
        else:
            if buf:
                chunks.append(buf)
            # hard-split very long paragraphs
            while len(p) > max_chars:
                cut = p.rfind(". ", 0, max_chars)
                cut = cut + 1 if cut > max_chars // 2 else max_chars
                chunks.append(p[:cut].strip())
                p = p[cut:].strip()
            buf = p
    if buf:
        chunks.append(buf)
    return chunks


def ingest_pdf(path: Path) -> list[Unit]:
    import fitz  # pymupdf

    units: list[Unit] = []
    doc = fitz.open(path)
    for pno, page in enumerate(doc, start=1):
        text = page.get_text("text").strip()
        n_images = len(page.get_images(full=True))
        if not text and n_images == 0:
            continue
        for ci, chunk in enumerate(_chunk_paragraphs(text)):
            if len(chunk) < 40:
                continue
            units.append(
                Unit(
                    id=_uid(path.name, f"p{pno}", ci),
                    text=chunk,
                    source_type="pdf",
                    source_file=path.name,
                    location=f"p.{pno}",
                    page=pno,
                    image_count=n_images if ci == 0 else 0,
                )
            )
    doc.close()
    return units


def ingest_pptx(path: Path) -> list[Unit]:
    from pptx import Presentation

    units: list[Unit] = []
    prs = Presentation(str(path))
    for sno, slide in enumerate(prs.slides, start=1):
        parts = []
        n_images = 0
        for shape in slide.shapes:
            if shape.has_text_frame:
                t = shape.text_frame.text.strip()
                if t:
                    parts.append(t)
            if shape.shape_type == 13:  # PICTURE
                n_images += 1
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        body = "\n".join(parts)
        if notes:
            body += f"\n\n[Speaker notes] {notes}"
        if not body.strip():
            continue
        units.append(
            Unit(
                id=_uid(path.name, f"slide{sno}", 0),
                text=body.strip(),
                source_type="pptx",
                source_file=path.name,
                location=f"slide {sno}",
                image_count=n_images,
            )
        )
    return units


def ingest_text(path: Path) -> list[Unit]:
    text = path.read_text(encoding="utf-8", errors="replace")
    units = []
    for ci, chunk in enumerate(_chunk_paragraphs(text)):
        if len(chunk) < 40:
            continue
        units.append(
            Unit(
                id=_uid(path.name, "text", ci),
                text=chunk,
                source_type="text",
                source_file=path.name,
                location=f"para {ci + 1}",
            )
        )
    return units


_TS_RE = re.compile(
    r"(?:(\d{1,2}):)?(\d{1,2}):(\d{2})[.,]\d{3}\s*-->"
)


def _ts_to_sec(ts: str) -> int:
    parts = ts.replace(",", ".").split(":")
    parts = [float(p) for p in parts]
    if len(parts) == 3:
        return int(parts[0] * 3600 + parts[1] * 60 + parts[2])
    if len(parts) == 2:
        return int(parts[0] * 60 + parts[1])
    return int(parts[0])


def ingest_transcript(path: Path) -> list[Unit]:
    """VTT or SRT transcripts -> chunks with start timestamps (windowed ~60s)."""
    raw = path.read_text(encoding="utf-8", errors="replace")
    lines = raw.splitlines()
    entries: list[tuple[int, str]] = []
    cur_ts = None
    for line in lines:
        line = line.strip()
        m = _TS_RE.search(line)
        if m:
            cur_ts = _ts_to_sec(line.split("-->")[0].strip())
            continue
        if line and cur_ts is not None and not line.isdigit() and line.upper() != "WEBVTT":
            entries.append((cur_ts, line))
            cur_ts = None
    if not entries:
        # plain transcript without timestamps -> fall back to paragraphs
        return ingest_text(path)
    units: list[Unit] = []
    buf_text: list[str] = []
    buf_start = entries[0][0]
    last_ts = entries[0][0]
    for ts, text in entries:
        buf_text.append(text)
        if ts - buf_start >= 60:  # ~60s window
            mm, ss = divmod(buf_start, 60)
            hh, mm = divmod(mm, 60)
            units.append(
                Unit(
                    id=_uid(path.name, f"t{buf_start}", 0),
                    text=" ".join(buf_text).strip(),
                    source_type="transcript",
                    source_file=path.name,
                    location=f"{hh:02d}:{mm:02d}:{ss:02d}",
                )
            )
            buf_text = []
            buf_start = ts
        last_ts = ts
    if buf_text:
        mm, ss = divmod(buf_start, 60)
        hh, mm = divmod(mm, 60)
        units.append(
            Unit(
                id=_uid(path.name, f"t{buf_start}", 0),
                text=" ".join(buf_text).strip(),
                source_type="transcript",
                source_file=path.name,
                location=f"{hh:02d}:{mm:02d}:{ss:02d}",
            )
        )
    return units


def ingest_file(path: Path) -> list[Unit]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return ingest_pdf(path)
    if suffix == ".pptx":
        return ingest_pptx(path)
    if suffix in (".vtt", ".srt"):
        return ingest_transcript(path)
    if suffix in (".txt", ".md"):
        # heuristic: timestamped transcript?
        head = path.read_text(encoding="utf-8", errors="replace")[:500]
        if "-->" in head or head.startswith("WEBVTT"):
            return ingest_transcript(path)
        return ingest_text(path)
    raise ValueError(f"Unsupported file type: {suffix} (supported: pdf, pptx, txt, md, vtt, srt)")


class KnowledgeBase:
    """Collection of units + topic tags, persisted as JSON."""

    def __init__(self, units: list[Unit] | None = None):
        self.units: list[Unit] = units or []

    def add_file(self, path: Path) -> int:
        new = ingest_file(path)
        existing = {(u.source_file, u.location) for u in self.units}
        added = [u for u in new if (u.source_file, u.location) not in existing]
        self.units.extend(added)
        return len(added)

    def to_json(self) -> str:
        return json.dumps([u.to_dict() for u in self.units], ensure_ascii=False, indent=1)

    @classmethod
    def from_json(cls, raw: str) -> "KnowledgeBase":
        return cls([Unit(**d) for d in json.loads(raw)])

    def save(self, path: Path) -> None:
        path.write_text(self.to_json(), encoding="utf-8")

    @classmethod
    def load(cls, path: Path) -> "KnowledgeBase":
        return cls.from_json(path.read_text(encoding="utf-8"))

    def summary(self) -> dict:
        files: dict[str, int] = {}
        for u in self.units:
            files[u.source_file] = files.get(u.source_file, 0) + 1
        return {"total_units": len(self.units), "files": files}

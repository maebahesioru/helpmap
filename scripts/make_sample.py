"""Generate the sample course slides (PPTX) and lecture transcript (VTT)."""
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

OUT = Path(__file__).parent.parent / "data" / "sample_course"

SLIDES = [
    ("Introduction to Machine Learning", "Lecture 1 — what learning from data means", [
        "Machine learning: algorithms that improve with experience",
        "Task T, experience E, performance measure P",
        "Three paradigms: supervised, unsupervised, reinforcement",
    ]),
    ("Supervised vs Unsupervised", "The two workhorses", [
        "Supervised: labelled (x, y) pairs — classification & regression",
        "Unsupervised: unlabelled data — clustering, dimensionality reduction",
        "Reinforcement: agent, environment, reward signal",
    ]),
    ("Linear Regression", "Fitting a line to data", [
        "Prediction: y_hat = w·x + b",
        "Loss: mean squared error (MSE)",
        "Closed form: normal equations w = (X^T X)^{-1} X^T y",
        "O(d^3) inversion — gradient descent often preferred",
    ]),
    ("Gradient Descent", "The optimization engine", [
        "theta := theta - alpha * grad L(theta)",
        "Learning rate alpha: too small = slow, too large = diverge",
        "SGD uses mini-batches for cheaper steps",
    ]),
    ("Overfitting & Regularization", "When memorizing beats learning", [
        "Overfitting: great on train, poor on new data",
        "L2 (ridge): lambda * ||w||^2 — smooth shrinkage",
        "L1 (lasso): lambda * ||w||_1 — sparse, feature selection",
        "Tune lambda on validation, report on test",
    ]),
]

prs = Presentation()
prs.slide_width = Inches(13.33)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

for title, subtitle, bullets in SLIDES:
    slide = prs.slides.add_slide(blank)
    tb = slide.shapes.add_textbox(Inches(0.7), Inches(0.5), Inches(12), Inches(1.4))
    tf = tb.text_frame
    tf.text = title
    tf.paragraphs[0].font.size = Pt(34)
    tf.paragraphs[0].font.bold = True
    p = tf.add_paragraph()
    p.text = subtitle
    p.font.size = Pt(17)

    body = slide.shapes.add_textbox(Inches(0.9), Inches(2.2), Inches(11.5), Inches(4.5))
    bf = body.text_frame
    for i, b in enumerate(bullets):
        para = bf.paragraphs[0] if i == 0 else bf.add_paragraph()
        para.text = "•  " + b
        para.font.size = Pt(20)
        para.space_after = Pt(14)

    # speaker notes
    notes = slide.notes_slide.notes_text_frame
    notes.text = "Explain: " + " / ".join(bullets)

prs.save(OUT / "lecture_slides.pptx")
print("slides:", len(SLIDES))

# ---------------------------------------------------------------- VTT
VTT = """WEBVTT

00:00:00.000 --> 00:00:22.000
Welcome to lecture one. Machine learning is the study of algorithms that improve their performance on a task through experience. Formally, a program learns from experience E with respect to task T and performance measure P if its performance at T improves with E.

00:00:22.000 --> 00:00:48.000
There are three main paradigms. Supervised learning uses labelled examples, input-output pairs, and learns a mapping that generalizes to unseen data. Think spam classification or house price prediction.

00:00:48.000 --> 00:01:15.000
Unsupervised learning works with unlabelled data and tries to discover structure: clusters, low-dimensional representations, or anomalies. Customer segmentation is the classic example.

00:01:15.000 --> 00:01:40.000
Reinforcement learning trains an agent that interacts with an environment. The agent observes states, takes actions, and receives rewards, and the goal is a policy that maximizes expected cumulative reward.

00:01:40.000 --> 00:02:10.000
Now linear regression. We model a continuous target y as y_hat equals w dot x plus b. Training means picking w and b to minimize a loss, most commonly the mean squared error. For MSE there is a closed-form solution, the normal equations: w equals X transpose X inverse, times X transpose y.

00:02:10.000 --> 00:02:40.000
But inverting X transpose X costs order d cubed, which is painful when you have many features, and it can be numerically unstable. That is why in practice we usually prefer gradient descent.

00:02:40.000 --> 00:03:15.000
Gradient descent is iterative. You start from an initial guess and repeatedly step in the direction of the negative gradient: theta becomes theta minus alpha times the gradient of the loss. Alpha is the learning rate and it is the most important hyperparameter. Too small and training crawls. Too large and you overshoot and diverge.

00:03:15.000 --> 00:03:45.000
Stochastic gradient descent approximates the gradient with a small random mini-batch. Each step is much cheaper, and convergence is often faster in practice, but the loss curve gets noisier.

00:03:45.000 --> 00:04:20.000
Finally, overfitting and regularization. A model overfits when it fits the training data very well but generalizes poorly. The remedy is to add a penalty on complexity. L2, or ridge, adds lambda times the squared norm of w, shrinking weights smoothly. L1, or lasso, adds lambda times the absolute norm, which pushes some weights exactly to zero and performs feature selection. Always tune lambda on a validation set, and report final numbers on the test set.
"""

(OUT / "lecture_transcript.vtt").write_text(VTT, encoding="utf-8")
print("vtt written")
